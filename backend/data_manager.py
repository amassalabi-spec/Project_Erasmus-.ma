import json
import re
import hashlib
import time
from pathlib import Path
from typing import List, Tuple
from langchain_community.document_loaders import CSVLoader, PyPDFLoader, TextLoader
from langchain_community.embeddings import HuggingFaceEmbeddings
from langchain_community.vectorstores import Chroma
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pptx import Presentation

from backend.config import LOGGER, SUPPORTED_EXTENSIONS, RAGConfig
from backend.utils import ensure_directory, extract_urls_from_text

# Ce module gère la découverte, le chargement, le nettoyage et l'indexation des documents dans la VDB Chroma.
# Il inclut des fonctions pour filtrer les chunks de texte bruyants, dédupliquer les chunks, et formater les documents pour les prompts du LLM.

NOISE_PATTERNS = [
    r"\binstruction\b", r"\bprompt\b", r"\bprompts\b", r"\bdo not\b", r"\bdon't\b",
    r"\bignore previous\b", r"\bsystem:|assistant:|human:|user:\b", r"\bplease ignore\b",
    r"\binstructions?:\b", r"\bto do\b", r"\btask\b", r"\bexample input\b", r"\bexample output\b",
]
NOISE_RE = re.compile("|".join(NOISE_PATTERNS), re.IGNORECASE)

def is_noise_chunk(text: str) -> bool:
    if not text: return False
    text = text.replace('\n', ' ').strip()
    if len(text) < 20: return True
    return bool(NOISE_RE.search(text))

def discover_data_files(data_dir):
    data_dir = Path(data_dir)
    if not data_dir.exists(): return []
    return sorted(p for p in data_dir.rglob("*") if p.is_file() and p.suffix.lower() in SUPPORTED_EXTENSIONS)

def load_single_file(file_path):
    file_path = Path(file_path)
    suffix = file_path.suffix.lower()
    
    if suffix in [".jsonl", ".json"]:
        docs = []
        with open(file_path, "r", encoding="utf-8") as f:
            for i, line in enumerate(f):
                line = line.strip()
                if not line: continue
                try:
                    data = json.loads(line)
                    content = data.get("texte", data.get("text", data.get("page_content", data.get("content", ""))))
                    metadata = data.get("metadata", {})
                    if not isinstance(metadata, dict): metadata = {}
                    for k, v in data.items():
                        if k not in ["texte", "text", "page_content", "content", "metadata"]:
                            metadata[k] = v
                    if content: docs.append(Document(page_content=content, metadata=metadata))
                except json.JSONDecodeError:
                    LOGGER.warning(f"Skipping invalid JSON line {i+1} in {file_path}")
        return docs

    if suffix == ".pdf": return PyPDFLoader(str(file_path)).load()
    if suffix == ".csv": return CSVLoader(file_path=str(file_path), encoding="utf-8").load()
    if suffix == ".pptx":
        prs = Presentation(str(file_path))
        docs = []
        for i, slide in enumerate(prs.slides, start=1):
            texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, 'text') and shape.text]
            page_text = "\n".join(texts).strip()
            if page_text: docs.append(Document(page_content=page_text, metadata={"page": i}))
        return docs
    return TextLoader(str(file_path), encoding="utf-8").load()

def conservative_clean_text(text: str) -> str:
    if not text: return text
    text = re.sub(r"```[\s\S]*?```", "\n", text)
    text = "\n".join(l for l in text.splitlines() if not re.match(r"^\s*(>>>|\$|PS>|\w+\s*:\s)", l))
    text = "\n".join(l for l in text.splitlines() if not NOISE_RE.search(l))
    text = re.sub(r"\n{2,}", "\n\n", text)
    text = re.sub(r"[ \t]{2,}", " ", text)
    return text.strip()

def load_documents(config: RAGConfig):
    files = discover_data_files(config.data_dir)
    if not files: raise FileNotFoundError("Aucun fichier compatible trouvé dans ./data")

    documents = []
    for file_path in files:
        relative_parts = [part.lower() for part in Path(file_path).relative_to(config.data_dir).parts]
        
        if "01_global" in relative_parts: scope, profile = "global", "all"
        elif "02_maroc" in relative_parts: scope, profile = "maroc", "all"
        elif "02_profils" in relative_parts and "etudiants" in relative_parts: scope, profile = "local", "etudiant"
        elif "02_profils" in relative_parts and "etablissements" in relative_parts: scope, profile = "local", "etablissement"
        elif "02_profils" in relative_parts and "associations" in relative_parts: scope, profile = "local", "association"
        else: scope, profile = "global", "all"

        for doc in load_single_file(file_path):
            metadata = dict(doc.metadata or {})
            cleaned_content = conservative_clean_text(doc.page_content or "")
            if not cleaned_content: continue
            
            if "scope" not in metadata: metadata["scope"] = scope
            if "profile" not in metadata: metadata["profile"] = profile
            metadata["source"] = file_path.as_posix()
            metadata["file_type"] = file_path.suffix.lower()
            
            doc.page_content = cleaned_content
            doc.metadata = metadata
            documents.append(doc)
    return documents

def build_or_load_vectorstore(config: RAGConfig, force_reindex=False, filter_noisy_chunks=False) -> Chroma:
    ensure_directory(config.data_dir, create=False)
    persist_dir = ensure_directory(config.persist_dir, create=True)
    embeddings = HuggingFaceEmbeddings(model_name=config.embedding_model_name)

    if force_reindex:
        persist_dir = Path(config.persist_dir) / f"reindex_{int(time.time())}"
        config.persist_dir = str(persist_dir)
        persist_dir = ensure_directory(persist_dir, create=True)

    if not force_reindex:
        try:
            vectorstore = Chroma(
                persist_directory=config.persist_dir,
                embedding_function=embeddings,
                collection_name="lumina",
            )
            if vectorstore._collection.count() > 0: return vectorstore
        except Exception as exc:
            LOGGER.warning("Rechargement Chroma impossible, reconstruction: %s", exc)

    documents = load_documents(config)
    
    txt_docs = [d for d in documents if d.metadata.get("file_type") == ".txt"]
    other_docs = [d for d in documents if d.metadata.get("file_type") != ".txt"]

    default_splitter = RecursiveCharacterTextSplitter(
        chunk_size=config.chunk_size,
        chunk_overlap=config.chunk_overlap,
        separators=["\n\n", "\n", " ", ""],
    )
    
    txt_splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=200,
        separators=["\n\n", "\n", " ", ""],
    )
    
    chunks = default_splitter.split_documents(other_docs) + txt_splitter.split_documents(txt_docs)

    if filter_noisy_chunks:
        try: chunks = [chunk for chunk in chunks if not is_noise_chunk(chunk.page_content or "")]
        except Exception as exc: LOGGER.debug("Filtrage echoue: %s", exc)

    def dedupe_chunks_by_hash(chunks):
        unique_chunks = {}
        for chunk in chunks:
            content = re.sub(r"\s+", " ", (chunk.page_content or "").strip())
            if content: unique_chunks[hashlib.md5(content.encode("utf-8")).hexdigest()] = chunk
        return list(unique_chunks.values())

    try:
        before_n = len(chunks)
        chunks = dedupe_chunks_by_hash(chunks)
        if len(chunks) != before_n:
            LOGGER.info("Deduplication: %d -> %d", before_n, len(chunks))
    except Exception as exc:
        LOGGER.debug("Deduplication echouee: %s", exc)

    vectorstore = Chroma.from_documents(documents=chunks, embedding=embeddings, persist_directory=config.persist_dir, collection_name="lumina")
    vectorstore.persist()
    return vectorstore

def format_docs_for_prompt_llm(docs: List[Document]) -> Tuple[str, List[str]]:
    blocks = []
    urls = []
    
    def sort_key(doc):
        meta = doc.metadata or {}
        try: page = int(meta.get("page", 0))
        except (ValueError, TypeError): page = 0
        return (str(meta.get("source", "")), page)
        
    sorted_docs = sorted(docs, key=sort_key)
    
    for index, doc in enumerate(sorted_docs, start=1):
        metadata = dict(doc.metadata or {})
        source = str(metadata.get("source", "source")).strip()
        short_source = Path(source).name if source else "Inconnue"
        page = metadata.get("page", "?")
        
        content = (doc.page_content or "").strip()
        if not content: continue
            
        doc_urls = extract_urls_from_text(content)
        for url in doc_urls:
            if url not in urls: urls.append(url)
                
        blocks.append(f"--- Document {index} : {short_source} (Page {page}) ---\n{content}")

    return "\n\n".join(blocks).strip(), urls
